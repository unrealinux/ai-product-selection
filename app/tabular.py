"""表格导入：把任意来源的商品表映射成 ``ProductIn``。

适用于走不了官方 API 的渠道 —— 例如 1688 商家后台 / 分销后台导出的商品表、
第三方数据服务导出的 Excel、手工整理的 CSV。

设计要点：

1. **列名自动识别**：内置常见中英文列名别名，两轮匹配（精确 → 包含）。
2. **列映射可存可复用**：映射方案按「源列名集合指纹」存库，
   同一份报表再次导入时自动套用，不必重配。
3. **单位不猜**：重量按列名里的 ``(g)`` / ``(kg)`` / ``(斤)`` 判定，
   识别不出就取默认值并在提示里说明；售价支持元 / 分。
4. **缺列可推导**：只有采购价时按加价倍数推导售价；
   只有销量时按对数映射推导热度（复用抖音那套口径）。
5. **来源写进 note**：明确记录哪些字段来自表格、哪些是推导、哪些是缺省值。
"""

from __future__ import annotations

import csv
import hashlib
import logging
import math
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Sequence

from .crawler import Source
from .models import ProductIn
from .scoring import clamp, log_scale

logger = logging.getLogger(__name__)

#: 销量达到该值 → 热度记 100（与抖音数据源保持同一口径）
SALES_FOR_MAX_HEAT = 1_000_000

#: 目标字段 → 常见列名别名。匹配前会去掉空白/下划线/括号等噪音并转小写
FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "title": (
        "商品标题", "商品名称", "产品名称", "宝贝标题", "标题", "品名", "商品",
        "title", "product_name", "productname", "goods_name", "name",
    ),
    "external_id": (
        "商品id", "商品编号", "商品编码", "货号", "编号", "款号",
        "offerid", "offer_id", "product_id", "productid", "skuid", "sku", "id",
    ),
    "category": (
        "商品类目", "类目名称", "一级类目", "二级类目", "类目", "分类", "行业",
        "category", "category_name", "cate", "categoryname",
    ),
    "url": (
        "商品链接", "详情链接", "商品url", "链接", "地址",
        "detail_url", "product_url", "url", "link",
    ),
    "price": (
        "售价", "零售价", "销售价", "建议售价", "终端价", "对外价",
        "sale_price", "selling_price", "retail_price", "price",
    ),
    "cost": (
        "成本价", "成本", "采购价", "批发价", "进货价", "供货价", "单价", "起批价",
        "purchase_price", "wholesale_price", "supply_price", "cost",
    ),
    "sales": (
        "历史销量", "30天成交", "30天销量", "成交笔数", "成交量", "已售", "销量",
        "sold_count", "sales_count", "sold", "sales", "volume",
    ),
    "heat": (
        "需求热度", "搜索指数", "热度", "人气",
        "heat", "hot", "popularity", "search_index",
    ),
    "competition": (
        "竞争指数", "竞争度", "竞争", "competition", "competition_index",
    ),
    "weight_kg": (
        "商品重量", "单件重量", "毛重", "净重", "重量",
        "weight_kg", "gross_weight", "net_weight", "weight",
    ),
    "repurchase": (
        "复购率", "复购潜力", "复购", "回头率", "回购率",
        "repurchase_rate", "repeat_rate", "repurchase",
    ),
    "compliance_risk": (
        "合规风险", "风险等级", "合规", "风险",
        "compliance_risk", "risk_level", "compliance", "risk",
    ),
    "virality": (
        "内容传播力", "传播潜力", "传播力", "传播",
        "virality", "spread", "spreadability",
    ),
    "note": (
        "备注", "说明", "描述", "商品描述", "remark", "comment", "description", "note",
    ),
}

#: 需要归一化到 0-100 的评分列。若整列最大值 ≤ 1，视为比例并乘以 100
RATE_FIELDS: tuple[str, ...] = (
    "repurchase", "compliance_risk", "virality", "heat", "competition",
)

#: 重量单位 → 千克换算系数
WEIGHT_UNIT_FACTORS: dict[str, float] = {
    "kg": 1.0, "g": 0.001, "mg": 0.000001, "jin": 0.5,
}
WEIGHT_UNIT_LABELS: dict[str, str] = {
    "kg": "千克", "g": "克", "mg": "毫克", "jin": "斤",
}

#: 售价单位 → 元换算系数
PRICE_UNIT_FACTORS: dict[str, float] = {"yuan": 1.0, "fen": 0.01}
PRICE_UNIT_LABELS: dict[str, str] = {"yuan": "元", "fen": "分"}

#: 支持的表头扩展名
CSV_SUFFIXES = {".csv", ".tsv", ".txt"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
LEGACY_EXCEL_SUFFIXES = {".xls"}

ENCODING_CANDIDATES = ("utf-8-sig", "utf-8", "gb18030", "big5", "latin-1")

_NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")
_COLUMN_NOISE_RE = re.compile(r"[\s_\-（）()\[\]【】{}:：*·、,，]")

#: 列名里的重量单位标记。用词边界避免把 "weight" 里的 g 误判成克
_WEIGHT_UNIT_PATTERNS: tuple[tuple[re.Pattern[str], str], ...] = (
    (re.compile(r"(mg|毫克)", re.I), "mg"),
    (re.compile(r"(kg|千克|公斤)", re.I), "kg"),
    (re.compile(r"(斤)", re.I), "jin"),
    (re.compile(r"(?<![a-zA-Z])(g|克)(?![a-zA-Z])", re.I), "g"),
)
_PRICE_UNIT_TOKENS: tuple[tuple[str, str], ...] = (
    ("yuan", "yuan"), ("cny", "yuan"), ("rmb", "yuan"), ("fen", "fen"),
)

MISSING_TOKENS = {"", "-", "--", "/", "n/a", "na", "null", "none", "nan", "暂无", "无"}

#: 目标字段 → 中文名，用于 CLI / UI 展示与手写映射
FIELD_LABELS: dict[str, str] = {
    "title": "商品标题",
    "external_id": "商品编号",
    "category": "类目",
    "url": "商品链接",
    "price": "售价",
    "cost": "成本",
    "sales": "销量",
    "heat": "热度",
    "competition": "竞争度",
    "weight_kg": "重量",
    "repurchase": "复购",
    "compliance_risk": "合规风险",
    "virality": "传播力",
    "note": "备注",
}


# --------------------------------------------------------------------------- #
# 列名处理
# --------------------------------------------------------------------------- #

def normalize_column(name: Any) -> str:
    """归一化列名用于匹配：去噪音字符并转小写。"""
    return _COLUMN_NOISE_RE.sub("", str(name or "")).lower()


def auto_map(columns: Sequence[str]) -> dict[str, str]:
    """按别名自动推断「目标字段 → 源列名」。两轮：精确匹配优先，其次包含匹配。"""
    lookup: dict[str, str] = {}
    for column in columns:
        lookup.setdefault(normalize_column(column), column)

    mapping: dict[str, str] = {}
    used: set[str] = set()

    for exact in (True, False):
        for target, aliases in FIELD_ALIASES.items():
            if target in mapping:
                continue
            for alias in aliases:
                key = normalize_column(alias)
                found = lookup.get(key) if exact else next(
                    (column for norm, column in lookup.items()
                     if key and key in norm and column not in used),
                    None,
                )
                if found and found not in used:
                    mapping[target] = found
                    used.add(found)
                    break
    return mapping


def detect_weight_unit(column: str) -> Optional[str]:
    """从列名推断重量单位（识别不出返回 None，由调用方取默认值）。"""
    for pattern, unit in _WEIGHT_UNIT_PATTERNS:
        if pattern.search(str(column or "")):
            return unit
    return None


def detect_price_unit(column: str) -> Optional[str]:
    """从列名推断金额单位（识别不出返回 None）。"""
    raw = str(column or "")
    normalized = normalize_column(raw)
    if "分" in raw and not re.search(r"分钟|分析|分享", raw):
        return "fen"
    if "元" in raw or "块" in raw:
        return "yuan"
    if any(token in normalized for token in ("yuan", "cny", "rmb")):
        return "yuan"
    if "fen" in normalized:
        return "fen"
    return None


def column_fingerprint(columns: Sequence[str]) -> str:
    """源列名集合的指纹，用于「同一份报表」的映射自动复用。"""
    normalized = sorted(normalize_column(column) for column in columns)
    return hashlib.sha1("|".join(normalized).encode("utf-8")).hexdigest()[:16]


def parse_mapping(text: str) -> dict[str, str]:
    """解析 ``"title=商品标题,cost=批发价(元)"`` 形式的列映射。

    左边可用英文字段名或中文名（如 ``商品标题=标题列``）。
    """
    if not text or not text.strip():
        return {}
    label_to_key = {label: key for key, label in FIELD_LABELS.items()}
    result: dict[str, str] = {}
    for chunk in text.replace("，", ",").split(","):
        chunk = chunk.strip()
        if not chunk:
            continue
        if "=" not in chunk:
            raise ValueError(f"映射片段缺少 '='：{chunk!r}")
        left, _, column = chunk.partition("=")
        left = left.strip()
        key = left if left in FIELD_LABELS else label_to_key.get(left)
        if key is None:
            raise ValueError(
                f"未知字段 {left!r}，可选：{', '.join(FIELD_LABELS)}"
            )
        column = column.strip()
        if not column:
            raise ValueError(f"{left} 对应的列名为空")
        result[key] = column
    return result


# --------------------------------------------------------------------------- #
# 取值解析
# --------------------------------------------------------------------------- #

def parse_number(value: Any) -> Optional[float]:
    """宽松解析数值：容忍千分位、百分号、货币符号与「暂无」这类占位符。"""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        return None if math.isnan(number) or math.isinf(number) else number
    text = str(value).strip()
    if text.lower() in MISSING_TOKENS:
        return None
    text = text.replace(",", "").replace("，", "").replace(" ", "")
    match = _NUMBER_RE.search(text)
    if not match:
        return None
    try:
        return float(match.group())
    except ValueError:
        return None


@dataclass
class RateScale:
    """评分列的量纲判断。

    真实的导出表经常混用两种写法（例如复购率列里既有 ``35%`` 又有 ``0.42``），
    所以**按值判定**而不是整列一个倍数：

    - 带百分号的值 → 本身就是 0-100 数轴上的数，原样使用
    - 不带百分号的值 → 看该列**这些值**的最大值：≤ 1 视为 0-1 比例（×100），
      否则视为 0-100（×1）
    """

    plain_factor: float = 1.0
    #: 同一列里既有 ≤1 又有 >1 的不带百分号的值 —— 量纲真实歧义，无法推断
    ambiguous: bool = False

    @classmethod
    def from_values(cls, raw_values: Sequence[Any]) -> "RateScale":
        plain: list[float] = []
        for value in raw_values:
            text = str(value).strip() if value is not None else ""
            if not text or "%" in text:
                continue
            number = parse_number(text)
            if number is not None:
                plain.append(number)
        if not plain:
            return cls()

        has_small = any(number <= 1.0 for number in plain)
        has_large = any(number > 1.0 for number in plain)
        if has_small and has_large:
            # 不猜：按 0-100 处理并让调用方提醒用户核对
            return cls(plain_factor=1.0, ambiguous=True)
        if has_large:
            return cls(plain_factor=1.0)
        return cls(plain_factor=100.0)

    def apply(self, raw_value: Any) -> Optional[float]:
        """返回归一化到 0-100 的值；无法解析返回 None。"""
        text = str(raw_value).strip() if raw_value is not None else ""
        number = parse_number(raw_value)
        if number is None:
            return None
        if "%" in text:
            return number
        return number * self.plain_factor

    @property
    def rescaled(self) -> bool:
        return self.plain_factor != 1.0


def rate_scale(raw_values: Sequence[Any]) -> float:
    """兼容旧接口：返回该评分列的放大倍数。"""
    return RateScale.from_values(raw_values).plain_factor


# --------------------------------------------------------------------------- #
# 读表
# --------------------------------------------------------------------------- #

@dataclass
class TableData:
    """读到的表格。"""

    columns: list[str]
    rows: list[dict[str, Any]]
    encoding: str = ""
    delimiter: str = ""
    sheet: str = ""
    path: str = ""

    def __len__(self) -> int:
        return len(self.rows)

    def describe(self) -> str:
        bits = [f"{len(self.rows)} 行 × {len(self.columns)} 列"]
        if self.encoding:
            bits.append(f"编码 {self.encoding}")
        if self.delimiter:
            bits.append(f"分隔符 {self.delimiter!r}")
        if self.sheet:
            bits.append(f"工作表 {self.sheet!r}")
        return "，".join(bits)


def _decode(raw: bytes) -> tuple[str, str]:
    """按候选编码解码，返回 ``(文本, 编码名)``。"""
    for encoding in ENCODING_CANDIDATES:
        try:
            return raw.decode(encoding), encoding
        except (UnicodeDecodeError, LookupError):
            continue
    # 兜底：替换非法字节，不抛错
    return raw.decode("utf-8", errors="replace"), "utf-8(replace)"


def _dedupe_columns(header: Sequence[Any]) -> list[str]:
    """清理表头：去空白、补空列名、去重。"""
    columns: list[str] = []
    seen: dict[str, int] = {}
    for index, cell in enumerate(header, start=1):
        name = str(cell).strip() if cell is not None else ""
        name = name or f"列{index}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        columns.append(name)
    return columns


def read_csv(path: Path | str, *, encoding: Optional[str] = None,
             delimiter: Optional[str] = None) -> TableData:
    """读取 CSV / TSV，自动识别编码与分隔符。"""
    path = Path(path)
    text, used_encoding = _decode(path.read_bytes())
    if encoding:
        text = path.read_text(encoding=encoding, errors="replace")
        used_encoding = encoding

    if delimiter is None:
        sample = "\n".join(text.splitlines()[:20])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = "\t" if "\t" in sample and "," not in sample else ","

    reader = csv.reader(text.splitlines(), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        return TableData([], [], encoding=used_encoding, delimiter=delimiter, path=str(path))

    columns = _dedupe_columns(header)
    rows: list[dict[str, Any]] = []
    for values in reader:
        if not any(str(v).strip() for v in values):
            continue  # 跳过空行
        rows.append({column: (values[i] if i < len(values) else "")
                     for i, column in enumerate(columns)})

    return TableData(columns, rows, encoding=used_encoding,
                     delimiter=delimiter, path=str(path))


def read_excel(path: Path | str, *, sheet: int | str = 0) -> TableData:
    """读取 xlsx / xlsm。``.xls``（旧二进制格式）不支持，会提示另存。"""
    path = Path(path)
    if path.suffix.lower() in LEGACY_EXCEL_SUFFIXES:
        raise ValueError(
            f"不支持旧版 .xls 格式（{path.name}）。"
            "请用 Excel 另存为 .xlsx，或导出为 CSV 后再导入。"
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "读取 Excel 需要 openpyxl，请先安装：pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet, str):
            if sheet not in workbook.sheetnames:
                raise ValueError(
                    f"工作表 {sheet!r} 不存在，可用：{', '.join(workbook.sheetnames)}"
                )
            worksheet = workbook[sheet]
        else:
            worksheet = workbook.worksheets[sheet]
        sheet_name = worksheet.title

        iterator = worksheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return TableData([], [], sheet=sheet_name, path=str(path))

        columns = _dedupe_columns(header)
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if values is None or not any(v is not None and str(v).strip() for v in values):
                continue
            rows.append({column: (values[i] if i < len(values) else None)
                         for i, column in enumerate(columns)})
    finally:
        workbook.close()

    return TableData(columns, rows, sheet=sheet_name, path=str(path))


def read_table(path: Path | str, **kwargs: Any) -> TableData:
    """按扩展名自动选择 CSV / Excel 读取器。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{path}")
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES or suffix in LEGACY_EXCEL_SUFFIXES:
        return read_excel(path, sheet=kwargs.get("sheet", 0))
    return read_csv(path, encoding=kwargs.get("encoding"), delimiter=kwargs.get("delimiter"))


def list_sheets(path: Path | str) -> list[str]:
    """列出 Excel 的工作表名；非 Excel 或依赖缺失时返回空列表。"""
    path = Path(path)
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except Exception:  # noqa: BLE001 - 仅用于 UI 提示，失败不致命
        return []


def is_excel(path: Path | str) -> bool:
    return Path(path).suffix.lower() in (EXCEL_SUFFIXES | LEGACY_EXCEL_SUFFIXES)


# --------------------------------------------------------------------------- #
# 列映射
# --------------------------------------------------------------------------- #

@dataclass
class ColumnMapping:
    """目标字段 → 源列名，外加单位与推导参数。"""

    fields: dict[str, str] = field(default_factory=dict)
    price_unit: str = "yuan"
    weight_unit: str = "kg"
    markup: float = 2.5
    source: str = "table"
    default_category: str = "未分类"

    def to_dict(self) -> dict[str, Any]:
        return {
            "fields": dict(self.fields),
            "price_unit": self.price_unit,
            "weight_unit": self.weight_unit,
            "markup": self.markup,
            "source": self.source,
            "default_category": self.default_category,
        }

    @classmethod
    def from_dict(cls, data: Mapping[str, Any]) -> "ColumnMapping":
        fields = dict(data.get("fields") or {})
        return cls(
            fields=fields,
            price_unit=str(data.get("price_unit") or "yuan"),
            weight_unit=str(data.get("weight_unit") or "kg"),
            markup=float(data.get("markup") or 2.5),
            source=str(data.get("source") or "table"),
            default_category=str(data.get("default_category") or "未分类"),
        )

    @classmethod
    def auto(cls, columns: Sequence[str], **overrides: Any) -> "ColumnMapping":
        """自动推断映射；重量/售价单位也从列名猜，猜不到用默认值。"""
        fields = auto_map(columns)
        params: dict[str, Any] = {"fields": fields}
        if "weight_kg" in fields:
            params["weight_unit"] = detect_weight_unit(fields["weight_kg"]) or "kg"
        if "price" in fields:
            params["price_unit"] = detect_price_unit(fields["price"]) or "yuan"
        if "cost" in fields:
            params["price_unit"] = detect_price_unit(fields["cost"]) or params.get(
                "price_unit", "yuan"
            )
        params.update({k: v for k, v in overrides.items() if v is not None})
        return cls(**params)

    def unmapped(self) -> list[str]:
        """返回未映射的必填/常用字段。"""
        return [name for name in ("title", "price", "cost", "sales", "weight_kg")
                if name not in self.fields]

    def describe(self) -> str:
        parts = [f"{target}←{column}" for target, column in sorted(self.fields.items())]
        return "，".join(parts) if parts else "（未映射任何列）"


# --------------------------------------------------------------------------- #
# 组装商品
# --------------------------------------------------------------------------- #

@dataclass
class BuildResult:
    """导入结果。"""

    products: list[ProductIn] = field(default_factory=list)
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)
    provenance: str = ""

    def summary(self) -> str:
        text = f"解析出 {len(self.products)} 个商品"
        if self.skipped:
            text += f"，跳过 {self.skipped} 行"
        return text + "。"


def build_provenance(mapping: ColumnMapping, columns: Sequence[str],
                     derived_notes: Sequence[str],
                     derived_targets: Iterable[str] = ()) -> str:
    """生成 note 尾部的数据来源说明。

    ``derived_targets`` 是那些**没有对应列、但被推导出来**的字段
    （例如由销量推出热度），它们不算「未提供」。
    """
    provided = set(mapping.fields) | set(derived_targets)
    bits = [f"来源：表格导入（{mapping.source}）"]
    if derived_notes:
        bits.append("推导：" + "、".join(derived_notes))
    missing = [_LABEL_BY_TARGET[target] for target in _PROVENANCE_TARGETS
               if target not in provided]
    if missing:
        bits.append("表中未提供（保持缺省）：" + "、".join(missing))

    unknown_columns = [column for column in columns
                       if column not in mapping.fields.values()]
    if unknown_columns:
        preview = "、".join(unknown_columns[:5]) + ("…" if len(unknown_columns) > 5 else "")
        bits.append(f"表内未使用列：{preview}")
    return "；".join(bits)


#: 需要向用户交代来源的维度
_PROVENANCE_TARGETS: tuple[str, ...] = (
    "heat", "competition", "weight_kg", "repurchase", "compliance_risk", "virality",
)

_LABEL_BY_TARGET = {
    "heat": "热度", "competition": "竞争度", "weight_kg": "重量",
    "repurchase": "复购", "compliance_risk": "合规", "virality": "传播力",
}


def build_products(data: TableData, mapping: ColumnMapping) -> BuildResult:
    """把表格行转换成 ``ProductIn`` 列表。"""
    result = BuildResult()
    if not data.columns:
        result.warnings.append("表格里没有读到任何列。")
        return result
    if "title" not in mapping.fields:
        result.warnings.append("未映射「商品标题」列，无法生成商品。请在映射里指定标题列。")
        return result

    missing_columns = [column for column in mapping.fields.values()
                       if column not in data.columns]
    if missing_columns:
        result.warnings.append(
            "映射引用了表里不存在的列：" + "、".join(missing_columns)
        )

    weight_factor = WEIGHT_UNIT_FACTORS.get(mapping.weight_unit, 1.0)
    price_factor = PRICE_UNIT_FACTORS.get(mapping.price_unit, 1.0)
    rate_scales = {
        target: RateScale.from_values([row.get(column) for row in data.rows])
        for target, column in mapping.fields.items() if target in RATE_FIELDS
    }

    derived_notes: list[str] = []
    derived_targets: list[str] = []
    if "price" not in mapping.fields and "cost" in mapping.fields:
        derived_notes.append(f"售价 = 采购价 × {mapping.markup:g}")
        derived_targets.append("price")
    if "heat" not in mapping.fields and "sales" in mapping.fields:
        derived_notes.append("热度 = 销量对数映射")
        derived_targets.append("heat")
    if mapping.weight_unit != "kg" and "weight_kg" in mapping.fields:
        derived_notes.append(
            f"重量按{WEIGHT_UNIT_LABELS.get(mapping.weight_unit, mapping.weight_unit)}换算"
        )
    if mapping.price_unit != "yuan":
        derived_notes.append(
            f"金额按{PRICE_UNIT_LABELS.get(mapping.price_unit, mapping.price_unit)}换算"
        )
    for target, scale in rate_scales.items():
        if scale.ambiguous:
            result.warnings.append(
                f"「{mapping.fields[target]}」列量纲不一致（同时存在 ≤1 与 >1 的值），"
                "已按 0-100 处理，请核对后再依赖排序。"
            )
        if scale.rescaled:
            derived_notes.append(f"{_LABEL_BY_TARGET.get(target, target)}按 0-1 比例×100")
        derived_targets.append(target)

    seen_ids: set[str] = set()
    for index, row in enumerate(data.rows, start=2):  # 第 1 行是表头
        def value(target: str) -> Any:
            column = mapping.fields.get(target)
            return row.get(column) if column else None

        title = str(value("title") or "").strip()
        if not title or title.lower() in MISSING_TOKENS:
            result.skipped += 1
            continue

        external_id = str(value("external_id") or "").strip()
        # 同一份表里重复出现的商品只保留第一条
        dedupe_key = external_id or title
        if dedupe_key in seen_ids:
            result.skipped += 1
            continue
        seen_ids.add(dedupe_key)

        cost_raw = parse_number(value("cost"))
        price_raw = parse_number(value("price"))
        if cost_raw is None and price_raw is None:
            result.skipped += 1
            if len(result.warnings) < 6:
                result.warnings.append(f"第 {index} 行缺少价格信息，已跳过（{title[:16]}）")
            continue

        if price_raw is not None:
            price = price_raw * price_factor
            cost = (cost_raw * price_factor) if cost_raw is not None else 0.0
        else:
            cost = (cost_raw or 0.0) * price_factor
            price = cost * max(0.0, mapping.markup)
        # 成本高于售价时按售价封顶，避免出现负毛利把分数打到 0 以下
        cost = min(cost, price)

        sales = parse_number(value("sales"))
        heat_scale = rate_scales.get("heat")
        heat_parsed = heat_scale.apply(value("heat")) if heat_scale else None
        heat = (clamp(heat_parsed) if heat_parsed is not None
                else log_scale(sales or 0, SALES_FOR_MAX_HEAT))

        weight_raw = parse_number(value("weight_kg"))
        weight = clamp(weight_raw * weight_factor, 0.0, 50.0) if weight_raw is not None else 0.5

        def rate(target: str, default: float) -> float:
            scale = rate_scales.get(target)
            raw = value(target)
            parsed = scale.apply(raw) if scale else parse_number(raw)
            return default if parsed is None else clamp(parsed)

        note_parts = [str(value("note") or "").strip()]
        if external_id:
            note_parts.insert(0, f"外部编号 {external_id}")
        if sales:
            note_parts.append(f"销量 {sales:g}")
        if cost_raw is not None and price_raw is None:
            note_parts.append(f"采购价 {cost:.2f} 元")
        note = " ".join(part for part in note_parts if part)

        try:
            product = ProductIn(
                title=title,
                external_id=external_id[:64],
                category=str(value("category") or "").strip()[:64] or mapping.default_category,
                price=round(price, 2),
                cost=round(cost, 2),
                source=mapping.source,
                url=str(value("url") or "").strip()[:500],
                heat=heat,
                competition=rate("competition", 50.0),
                weight_kg=round(weight, 4) if weight > 0 else 0.5,
                repurchase=rate("repurchase", 50.0),
                compliance_risk=rate("compliance_risk", 20.0),
                virality=rate("virality", 50.0),
                note=note,
            )
        except Exception as exc:  # noqa: BLE001 - 单行脏数据不应中断整批导入
            result.skipped += 1
            if len(result.warnings) < 6:
                result.warnings.append(f"第 {index} 行校验失败，已跳过（{title[:16]}）：{exc}")
            continue

        result.products.append(product)

    result.provenance = build_provenance(
        mapping, data.columns, derived_notes, derived_targets
    )
    return result


def apply_provenance(products: Sequence[ProductIn], provenance: str) -> list[ProductIn]:
    """把数据来源说明写进每个商品的 note。"""
    from .enrich import PROVENANCE_SEP

    if not provenance:
        return list(products)
    return [
        product.model_copy(update={"note": f"{product.note}{PROVENANCE_SEP}{provenance}"})
        for product in products
    ]


# --------------------------------------------------------------------------- #
# 数据源
# --------------------------------------------------------------------------- #

class TabularSource(Source):
    """表格数据源：读文件 → 映射 → ``ProductIn``。"""

    name = "table"

    def __init__(
        self,
        path: Path | str,
        mapping: Optional[ColumnMapping | Mapping[str, Any]] = None,
        *,
        sheet: int | str = 0,
        encoding: Optional[str] = None,
        delimiter: Optional[str] = None,
        markup: float = 2.5,
        source_label: str = "",
        weight_unit: Optional[str] = None,
        price_unit: Optional[str] = None,
    ) -> None:
        self.path = Path(path)
        self.sheet = sheet
        self.encoding = encoding
        self.delimiter = delimiter
        self.markup = markup
        # 用文件名做来源标签，便于在榜单里按渠道筛选
        self.source_label = source_label or self.path.stem or "table"
        self._weight_unit = weight_unit
        self._price_unit = price_unit
        self._mapping = mapping
        self.last_result: Optional[BuildResult] = None
        self.last_data: Optional[TableData] = None

    def resolve_mapping(self, columns: Sequence[str]) -> ColumnMapping:
        """确定使用的映射：显式传入 > 自动推断。"""
        if isinstance(self._mapping, ColumnMapping):
            return self._mapping
        if isinstance(self._mapping, Mapping):
            return ColumnMapping.from_dict(self._mapping)
        return ColumnMapping.auto(
            columns,
            markup=self.markup,
            source=self.source_label,
            weight_unit=self._weight_unit,
            price_unit=self._price_unit,
        )

    def load(self) -> TableData:
        """读表（结果缓存，便于 UI 先预览再导入）。"""
        if self.last_data is None:
            self.last_data = read_table(
                self.path, sheet=self.sheet,
                encoding=self.encoding, delimiter=self.delimiter,
            )
        return self.last_data

    def fetch(self) -> list[ProductIn]:
        data = self.load()
        mapping = self.resolve_mapping(data.columns)
        result = build_products(data, mapping)
        self.last_result = result
        if result.warnings:
            for warning in result.warnings:
                logger.warning("表格导入：%s", warning)
        return apply_provenance(result.products, result.provenance)
